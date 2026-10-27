"""
text_extractor.py
-----------------
Multi-format document text extraction service:
1. Plain text (.txt)
2. PDF documents via pdfplumber
3. Word documents (.docx) via python-docx
4. Images and scanned documents via Tesseract OCR (pytesseract)
Provides a friendly fallback message if OCR is not installed on the system.
"""

import io
from pathlib import Path
from typing import Dict, Any, Optional

# Optional library imports with safe detection
try:
    import pdfplumber
    PDFPLUMBER_AVAILABLE = True
except ImportError:
    PDFPLUMBER_AVAILABLE = False

try:
    import docx
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    from PIL import Image
    import pytesseract
    PYTESSERACT_AVAILABLE = True
except ImportError:
    PYTESSERACT_AVAILABLE = False


def extract_from_plain_text(file_bytes: bytes) -> str:
    """Extracts text from UTF-8 / Latin-1 encoded plain text files."""
    try:
        return file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        return file_bytes.decode("latin-1", errors="ignore")


def extract_from_pdf(file_bytes: bytes) -> Dict[str, Any]:
    """Extracts text from PDF documents page by page using pdfplumber."""
    if not PDFPLUMBER_AVAILABLE:
        return {
            "text": "",
            "pages": 0,
            "warning": "pdfplumber library is not installed. Please paste the text directly."
        }

    extracted_pages = []
    try:
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            page_count = len(pdf.pages)
            for page_num, page in enumerate(pdf.pages, start=1):
                page_text = page.extract_text()
                if page_text:
                    extracted_pages.append(page_text.strip())

        full_text = "\n\n".join(extracted_pages).strip()
        
        # If PDF has pages but 0 text, it might be a scanned image PDF
        if page_count > 0 and not full_text:
            return {
                "text": "",
                "pages": page_count,
                "is_scanned": True,
                "warning": "This PDF appears to be a scanned image without a selectable text layer. Please use OCR or paste the text."
            }

        return {
            "text": full_text,
            "pages": page_count,
            "is_scanned": False,
            "warning": None
        }
    except Exception as e:
        return {
            "text": "",
            "pages": 0,
            "warning": f"Could not read PDF file: {str(e)}"
        }


def extract_from_docx(file_bytes: bytes) -> Dict[str, Any]:
    """Extracts text from Microsoft Word .docx files using python-docx."""
    if not DOCX_AVAILABLE:
        return {
            "text": "",
            "warning": "python-docx library is not installed. Please paste the text directly."
        }

    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        
        # Also extract text inside tables if any
        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_text:
                    table_texts.append(row_text)

        all_text = "\n".join(paragraphs)
        if table_texts:
            all_text += "\n\nTable Data:\n" + "\n".join(table_texts)

        return {
            "text": all_text.strip(),
            "paragraph_count": len(paragraphs),
            "warning": None
        }
    except Exception as e:
        return {
            "text": "",
            "warning": f"Could not read Word DOCX file: {str(e)}"
        }


def extract_from_image(file_bytes: bytes) -> Dict[str, Any]:
    """
    Extracts text from image files (PNG, JPG, JPEG) using pytesseract OCR.
    If Tesseract OCR engine is not installed on the system, returns a friendly message.
    """
    if not PYTESSERACT_AVAILABLE:
        return {
            "text": "",
            "ocr_available": False,
            "warning": "Image OCR requires 'pytesseract' and 'Pillow'. Please paste the offer letter text directly."
        }

    try:
        image = Image.open(io.BytesIO(file_bytes))
        # Attempt OCR
        text = pytesseract.image_to_string(image)
        return {
            "text": text.strip(),
            "ocr_available": True,
            "warning": None
        }
    except pytesseract.TesseractNotFoundError:
        return {
            "text": "",
            "ocr_available": False,
            "warning": (
                "Tesseract OCR engine is not installed on this system. "
                "No worries! You can simply copy and paste the text into the text box above to check it instantly."
            )
        }
    except Exception as e:
        return {
            "text": "",
            "ocr_available": False,
            "warning": f"Could not process image: {str(e)}. Please paste the text instead."
        }


def extract_document_text(filename: str, file_bytes: bytes) -> Dict[str, Any]:
    """
    Master extractor function: Detects file type from filename extension,
    routes to the appropriate extractor, and returns extracted text and metadata.
    """
    if not file_bytes:
        return {
            "success": False,
            "text": "",
            "method": "None",
            "filename": filename,
            "warning": "Uploaded file is empty."
        }

    ext = Path(filename).suffix.lower()

    if ext in [".txt", ".log", ".md"]:
        text = extract_from_plain_text(file_bytes)
        return {
            "success": True,
            "text": text,
            "method": "Plain Text",
            "filename": filename,
            "warning": None
        }

    elif ext == ".pdf":
        pdf_res = extract_from_pdf(file_bytes)
        return {
            "success": bool(pdf_res["text"]),
            "text": pdf_res["text"],
            "method": "PDF Plumber",
            "filename": filename,
            "pages": pdf_res.get("pages", 0),
            "warning": pdf_res.get("warning")
        }

    elif ext in [".docx", ".doc"]:
        docx_res = extract_from_docx(file_bytes)
        return {
            "success": bool(docx_res["text"]),
            "text": docx_res["text"],
            "method": "Word DOCX",
            "filename": filename,
            "warning": docx_res.get("warning")
        }

    elif ext in [".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff"]:
        img_res = extract_from_image(file_bytes)
        return {
            "success": bool(img_res["text"]),
            "text": img_res["text"],
            "method": "Tesseract OCR",
            "filename": filename,
            "warning": img_res.get("warning")
        }

    else:
        # Attempt plain text read as fallback
        fallback_text = extract_from_plain_text(file_bytes)
        return {
            "success": bool(fallback_text.strip()),
            "text": fallback_text,
            "method": "Generic Fallback",
            "filename": filename,
            "warning": f"Unsupported file extension '{ext}'. Processed as plain text."
        }
