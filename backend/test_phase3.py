"""
test_phase3.py
--------------
Automated test suite for Phase 3:
1. Multi-format text extractor (TXT, PDF, DOCX, Image OCR fallback)
2. Scikit-learn Machine Learning Classifier (TF-IDF + Logistic Regression)
3. Multimodal Explainable Scoring Engine (Trust Score, Verdict, Red Banners)
"""

import sys
import unittest
import io
from pathlib import Path

# Add backend directory to system path
sys.path.insert(0, str(Path(__file__).resolve().parent))

from services.text_extractor import (
    extract_from_plain_text,
    extract_from_docx,
    extract_from_image,
    extract_document_text
)
from services.ml_model import ScamClassifier, scam_classifier
from services.scoring import calculate_comprehensive_trust_score


class TestPhase3MLAndScoring(unittest.TestCase):

    def test_01_text_extractor_plain_text(self):
        """Tests extraction from raw text bytes."""
        sample_bytes = "Job Title: Software Engineer\nCompany: TCS\nLocation: Bangalore".encode("utf-8")
        extracted = extract_from_plain_text(sample_bytes)
        self.assertIn("Software Engineer", extracted)
        self.assertIn("TCS", extracted)
        print(" [PASS] 01. Text Extractor (Plain Text) Verified")

    def test_02_text_extractor_docx(self):
        """Tests extraction from Microsoft Word docx format."""
        import docx
        doc = docx.Document()
        doc.add_heading("Letter of Intent", 0)
        doc.add_paragraph("We are pleased to offer you the role of Systems Analyst.")
        
        bio = io.BytesIO()
        doc.save(bio)
        file_bytes = bio.getvalue()

        result = extract_from_docx(file_bytes)
        self.assertIn("Letter of Intent", result["text"])
        self.assertIn("Systems Analyst", result["text"])
        print(" [PASS] 02. Text Extractor (Word DOCX) Verified")

    def test_03_text_extractor_image_fallback(self):
        """Tests that image processing returns clean text or a friendly fallback without crashing."""
        from PIL import Image
        img = Image.new("RGB", (200, 100), color=(255, 255, 255))
        img_byte_arr = io.BytesIO()
        img.save(img_byte_arr, format='PNG')
        img_bytes = img_byte_arr.getvalue()

        res = extract_from_image(img_bytes)
        # Should return a dict with text and warning/ocr status without throwing an exception
        self.assertIsInstance(res, dict)
        self.assertIn("text", res)
        print(" [PASS] 03. Text Extractor (Image OCR & Friendly Fallback) Verified")

    def test_04_ml_classifier_training(self):
        """Tests that the ML pipeline is trained on the balanced dataset."""
        info = scam_classifier.get_info()
        self.assertTrue(info["is_trained"])
        self.assertGreaterEqual(info["train_samples_count"], 40)
        self.assertIn("Logistic Regression", info["engine"])
        print(f" [PASS] 04. ML Classifier Training Verified ({info['train_samples_count']} samples, {info['dataset_source']})")

    def test_05_ml_classifier_predictions(self):
        """Tests ML predictions on genuine vs scam texts."""
        scam_text = "Urgent selection in WFH Data Entry. Transfer refundable registration fee of Rs. 1,499 via PhonePe to activate work kit."
        scam_pred = scam_classifier.predict(scam_text)
        self.assertTrue(scam_pred["is_scam"])
        self.assertGreaterEqual(scam_pred["scam_probability"], 0.60)
        self.assertEqual(scam_pred["verdict"], "High scam probability")

        genuine_text = "Tata Consultancy Services Limited is pleased to offer you the position of Assistant System Engineer with CTC 3.85 LPA. TCS follows a zero-fee recruitment policy."
        gen_pred = scam_classifier.predict(genuine_text)
        self.assertFalse(gen_pred["is_scam"])
        self.assertLessEqual(gen_pred["scam_probability"], 0.40)
        self.assertEqual(gen_pred["verdict"], "Safe / Likely genuine")
        print(f" [PASS] 05. ML Classification Predictions Verified (Scam: {scam_pred['scam_probability']}, Genuine: {gen_pred['scam_probability']})")

    def test_06_scoring_genuine_offer(self):
        """Tests multimodal explainable trust score on a genuine TCS offer letter."""
        samples_dir = Path(__file__).resolve().parent / "samples"
        with open(samples_dir / "genuine_offer_letter.txt", "r", encoding="utf-8") as f:
            text = f.read()

        res = calculate_comprehensive_trust_score(text=text)
        self.assertGreaterEqual(res["trust_score"], 70)
        self.assertEqual(res["verdict"], "Safe / Likely Genuine")
        self.assertIsNone(res["red_banner_warning"])
        
        # Verify explainable positive reasons
        titles = [r["title"] for r in res["reasons"]]
        self.assertIn("Zero-Fee Policy Confirmed", titles)
        print(f" [PASS] 06. Genuine Offer Letter Trust Score Verified: {res['trust_score']}/100 ({res['verdict']})")

    def test_07_scoring_scam_offer(self):
        """Tests multimodal explainable trust score on a scam offer letter."""
        samples_dir = Path(__file__).resolve().parent / "samples"
        with open(samples_dir / "scam_offer_letter.txt", "r", encoding="utf-8") as f:
            text = f.read()

        res = calculate_comprehensive_trust_score(text=text)
        self.assertLessEqual(res["trust_score"], 25)
        self.assertEqual(res["verdict"], "High scam probability")
        self.assertIsNotNone(res["red_banner_warning"])
        self.assertIn("demands money", res["red_banner_warning"].lower())

        # Verify explainable critical reasons
        titles = [r["title"] for r in res["reasons"]]
        self.assertIn("Payment / Security Deposit Demanded", titles)
        print(f" [PASS] 07. Scam Offer Letter Trust Score Verified: {res['trust_score']}/100 ({res['verdict']})")

    def test_08_scoring_suspicious_offer(self):
        """Tests multimodal explainable trust score on a suspicious offer letter."""
        samples_dir = Path(__file__).resolve().parent / "samples"
        with open(samples_dir / "suspicious_offer_letter.txt", "r", encoding="utf-8") as f:
            text = f.read()

        res = calculate_comprehensive_trust_score(text=text)
        # Should be flagged as suspicious or low trust due to free gmail and vendor link fee
        self.assertLess(res["trust_score"], 70)
        print(f" [PASS] 08. Suspicious Offer Letter Trust Score Verified: {res['trust_score']}/100 ({res['verdict']})")


if __name__ == "__main__":
    print("\n" + "=" * 65)
    print(" [JobShield] RUNNING PHASE 3 TEST SUITE: ML & Scoring")
    print("=" * 65 + "\n")
    unittest.main()
